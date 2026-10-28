
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import LineChart, BarChart, Reference

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "cleaned"
OUTPUT = ROOT / "excel" / "Banking_Transaction_Analytics.xlsx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

# Load the five validated CSV exports.
files = {
    "Monthly Summary": ("monthly_summary.csv", ["transaction_month"]),
    "Channel Summary": ("channel_summary.csv", []),
    "Status Summary": ("status_summary.csv", []),
    "Data Quality": ("data_quality_summary.csv", []),
    "Cleaned Transactions": (
        "cleaned_transactions.csv",
        ["transaction_date"],
    ),
}

frames = {}
for sheet_name, (filename, date_columns) in files.items():
    path = DATA / filename
    if not path.exists():
        raise FileNotFoundError(f"Missing export: {path}")

    df = pd.read_csv(path, parse_dates=date_columns)
    frames[sheet_name] = df

wb = Workbook()
dashboard = wb.active
dashboard.title = "Dashboard"

# Professional dashboard header.
dashboard.merge_cells("A1:H1")
dashboard["A1"] = "BANKING TRANSACTION ANALYTICS"
dashboard["A1"].font = Font(size=20, bold=True, color="FFFFFF")
dashboard["A1"].fill = PatternFill("solid", fgColor="163A5F")
dashboard["A1"].alignment = Alignment(horizontal="center")
dashboard.row_dimensions[1].height = 34

dashboard.merge_cells("A2:H2")
dashboard["A2"] = (
    "Synthetic portfolio dataset | January–December 2025 | "
    "PostgreSQL + Python + Excel"
)
dashboard["A2"].font = Font(italic=True, color="44546A")
dashboard["A2"].alignment = Alignment(horizontal="center", wrap_text=True)

# KPI cards use formulas linked to the transaction sheet.
cards = [
    ("A4:B4", "A5:B6", "CLEANED TRANSACTIONS",
     "=COUNTA('Cleaned Transactions'!A2:A5001)", "#,##0"),
    ("C4:D4", "C5:D6", "SUCCESS RATE",
     '=IFERROR(COUNTIF(\'Cleaned Transactions\'!H2:H5001,"SUCCESS")/'
     'COUNTA(\'Cleaned Transactions\'!A2:A5001),0)', "0.0%"),
    ("E4:F4", "E5:F6", "AMOUNT QUALITY ISSUES",
     '=COUNTIF(\'Cleaned Transactions\'!K2:K5001,"<>VALID")', "#,##0"),
    ("G4:H4", "G5:H6", "MISSING CUSTOMER IDs",
     '=COUNTIF(\'Cleaned Transactions\'!L2:L5001,"MISSING")', "#,##0"),
]

for label_range, value_range, label, formula, number_format in cards:
    dashboard.merge_cells(label_range)
    dashboard.merge_cells(value_range)

    label_cell = dashboard[label_range.split(":")[0]]
    value_cell = dashboard[value_range.split(":")[0]]

    label_cell.value = label
    label_cell.font = Font(bold=True, color="FFFFFF", size=10)
    label_cell.fill = PatternFill("solid", fgColor="2F75A5")
    label_cell.alignment = Alignment(horizontal="center", vertical="center")

    value_cell.value = formula
    value_cell.font = Font(bold=True, size=18, color="163A5F")
    value_cell.fill = PatternFill("solid", fgColor="DDEBF7")
    value_cell.alignment = Alignment(horizontal="center", vertical="center")
    value_cell.number_format = number_format

dashboard.row_dimensions[4].height = 24
dashboard.row_dimensions[5].height = 25
dashboard.row_dimensions[6].height = 25

# Add data tabs with Excel tables and readable formatting.
table_names = {
    "Monthly Summary": "TblMonthlySummary",
    "Channel Summary": "TblChannelSummary",
    "Status Summary": "TblStatusSummary",
    "Data Quality": "TblDataQuality",
    "Cleaned Transactions": "TblCleanedTransactions",
}

for sheet_name, df in frames.items():
    ws = wb.create_sheet(sheet_name)
    ws.append(list(df.columns))

    for record in df.itertuples(index=False, name=None):
        values = []
        for value in record:
            if pd.isna(value):
                values.append(None)
            elif isinstance(value, pd.Timestamp):
                values.append(value.to_pydatetime())
            elif hasattr(value, "item"):
                values.append(value.item())
            else:
                values.append(value)
        ws.append(values)

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    # Create a filterable, styled Excel table.
    if ws.max_row >= 2:
        table = Table(
            displayName=table_names[sheet_name],
            ref=ws.dimensions,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showRowStripes=True,
            showColumnStripes=False,
            showFirstColumn=False,
            showLastColumn=False,
        )
        ws.add_table(table)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Set practical column widths.
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        header = str(column_cells[0].value or "")
        max_length = max(
            (len(str(c.value)) for c in column_cells[:100]),
            default=len(header),
        )
        ws.column_dimensions[letter].width = min(
            max(max_length + 2, 14), 28
        )

    # Date and monetary number formats.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            header = str(ws.cell(1, cell.column).value).lower()
            if header == "transaction_date" and cell.value:
                cell.number_format = "yyyy-mm-dd"
            elif header == "transaction_month" and cell.value:
                cell.number_format = "mmm yyyy"
            elif "amount_inr" in header or header == "amount":
                cell.number_format = '#,##0.00'

# Monthly trend chart.
monthly = wb["Monthly Summary"]
line = LineChart()
line.title = "Monthly Transaction Volume"
line.y_axis.title = "Transactions"
line.x_axis.title = "Month"
line.height = 7
line.width = 14
line.add_data(
    Reference(monthly, min_col=2, max_col=3, min_row=1, max_row=13),
    titles_from_data=True,
)
line.set_categories(
    Reference(monthly, min_col=1, min_row=2, max_row=13)
)
line.style = 13
dashboard.add_chart(line, "A9")

# Channel transaction volume chart.
channel = wb["Channel Summary"]
bar = BarChart()
bar.type = "bar"
bar.style = 10
bar.title = "Transactions by Channel"
bar.x_axis.title = "Transactions"
bar.y_axis.title = "Channel"
bar.height = 7
bar.width = 13
bar.add_data(
    Reference(channel, min_col=2, min_row=1, max_row=channel.max_row),
    titles_from_data=True,
)
bar.set_categories(
    Reference(channel, min_col=1, min_row=2, max_row=channel.max_row)
)
dashboard.add_chart(bar, "E9")

# Transaction status chart.
status = wb["Status Summary"]
status_chart = BarChart()
status_chart.type = "col"
status_chart.style = 11
status_chart.title = "Transactions by Status"
status_chart.y_axis.title = "Transactions"
status_chart.x_axis.title = "Status"
status_chart.height = 7
status_chart.width = 13
status_chart.add_data(
    Reference(status, min_col=2, min_row=1, max_row=status.max_row),
    titles_from_data=True,
)
status_chart.set_categories(
    Reference(status, min_col=1, min_row=2, max_row=status.max_row)
)
dashboard.add_chart(status_chart, "A24")

dashboard.merge_cells("E24:H27")
dashboard["E24"] = (
    "DATA NOTE\n"
    "This portfolio uses synthetic transactions. Statuses were randomly "
    "generated, so success rates and transaction values are illustrative "
    "and are not real banking benchmarks."
)
dashboard["E24"].alignment = Alignment(
    wrap_text=True, vertical="center"
)
dashboard["E24"].font = Font(italic=True, color="7F6000")
dashboard["E24"].fill = PatternFill("solid", fgColor="FFF2CC")

for col in "ABCDEFGH":
    dashboard.column_dimensions[col].width = 16

dashboard.freeze_panes = "A4"
dashboard.sheet_view.showGridLines = False

# Make Excel recalculate KPI formulas when the workbook opens.
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

wb.save(OUTPUT)

print("Excel workbook created successfully!")
print(f"File: {OUTPUT}")
print(f"Sheets: {', '.join(wb.sheetnames)}")
print("Charts: monthly trend, channel volume, transaction status")
